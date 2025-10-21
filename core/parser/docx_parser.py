from core.parser.base_parser import BaseParser
from langchain.schema import Document
from typing import List
import os
from docx import Document as DocxDocument


class DocxLoader(BaseParser):
    """加载DOCX文件内容的加载器"""

    def __init__(self, file_path: str):
        """
        初始化DOCX加载器

        Args:
            file_path: DOCX文件的路径
        """
        self.file_path = file_path
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"文件不存在: {file_path}")
        if not file_path.lower().endswith('.docx'):
            raise ValueError(f"不支持的文件格式，需要.docx: {file_path}")

    def load(self) -> List[Document]:
        """
        加载并解析DOCX文件内容

        Returns:
            包含DOCX内容的Document对象列表
        """
        try:
            doc = DocxDocument(self.file_path)
            paragraphs = []

            # 提取段落内容
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)

            # 提取表格内容
            for table_idx, table in enumerate(doc.tables, 1):
                table_content = []
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells]
                    if any(row_text):  # 只添加非空行
                        table_content.append("| ".join(row_text))

                if table_content:
                    paragraphs.append(f"\n表格 {table_idx}:\n" + "\n".join(table_content))

            content = "\n\n".join(paragraphs)
            metadata = {
                "source": self.file_path,
                "paragraph_count": len(doc.paragraphs),
                "table_count": len(doc.tables),
                "file_type": "docx"
            }

            return [Document(page_content=content, metadata=metadata)]

        except Exception as e:
            raise RuntimeError(f"解析DOCX文件失败: {str(e)}") from e
