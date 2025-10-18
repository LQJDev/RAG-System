import os.path
from core.loaders.base_loader import BaseLoader
from langchain.schema import Document
from typing import List
from pptx import Presentation

class PptLoader(BaseLoader):

    """加载PPT(.pptx)文件内容的加载器"""
    def __init__(self, file_path: str):
        """
        初始化PPT加载器

        Args:
            file_path: PPT文件的路径
        """
        self.file_path = file_path
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"文件不存在: {file_path}")
        if not file_path.lower().endswith('.pptx'):
            raise ValueError(f"不支持的文件格式，需要.pptx: {file_path}")


    def load(self) -> List[Document]:
        """
        加载并解析PPT文件内容

        Returns:
            包含PPT内容的Document对象列表
        """
        try:
            prs = Presentation(self.file_path)
            slides_content = []

            for slide_idx, slide in enumerate(prs.slides, 1):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text)

                if slide_text:
                    content = "\n\n".join(slide_text)
                    metadata = {
                        "source": self.file_path,
                        "slide_number": slide_idx,
                        "file_type": "pptx"
                    }
                    slides_content.append(Document(page_content=content, metadata=metadata))

            return slides_content

        except Exception as e:
            raise RuntimeError(f"解析PPT文件失败: {str(e)}") from e
